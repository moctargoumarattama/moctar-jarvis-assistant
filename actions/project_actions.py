import csv
import os
import subprocess
from pathlib import Path

import config


class ProjectAssistant:
    def __init__(self, projects=None):
        self.projects = projects or config.PROJECTS
        self.current_project_key = None

    def _resolve_project(self, project_key=None):
        key = project_key or self.current_project_key
        if not key:
            return None
        return self.projects.get(key)

    def _set_current_project(self, project_key):
        self.current_project_key = project_key

    def open_project(self, project_key, open_in_vscode=False):
        project = self._resolve_project(project_key)
        if not project:
            return "Je ne connais pas encore ce projet."
        if not project.path.exists():
            return f"Le chemin du projet {project.label} est introuvable. Mets-le a jour dans config.py."

        self._set_current_project(project_key)

        if open_in_vscode:
            return self.open_project_in_vscode(project_key)

        try:
            os.startfile(str(project.path))
            return f"J'ouvre le projet {project.label}."
        except Exception as exc:
            return f"Impossible d'ouvrir le projet {project.label} : {exc}"

    def open_project_in_vscode(self, project_key=None):
        project = self._resolve_project(project_key)
        if not project:
            return "Choisis d'abord un projet a ouvrir."
        if not project.path.exists():
            return f"Le chemin du projet {project.label} est introuvable. Mets-le a jour dans config.py."

        self._set_current_project(project.key)

        try:
            subprocess.Popen(
                ["cmd", "/c", "start", "", "code", str(project.path)],
                shell=False,
            )
            return f"J'ouvre {project.label} dans VS Code."
        except Exception as exc:
            return f"Impossible d'ouvrir VS Code pour {project.label} : {exc}"

    def launch_project_server(self, target=None):
        project = self._resolve_project(target if target in self.projects else None)
        if project is None:
            project = self._resolve_project()
        if project is None:
            framework = target if target else "ce framework"
            return f"Choisis d'abord un projet avant de lancer {framework}."
        if not project.path.exists():
            return f"Le chemin du projet {project.label} est introuvable. Mets-le a jour dans config.py."

        if target in {"flask", "fastapi"} and project.framework and project.framework != target:
            return f"Le projet actif {project.label} est configure en {project.framework}, pas en {target}."

        if not project.server_command:
            return f"Aucune commande serveur n'est configuree pour {project.label}."

        try:
            subprocess.Popen(
                list(project.server_command),
                cwd=project.server_cwd or project.path,
                shell=False,
            )
            self._set_current_project(project.key)
            framework = project.framework or "serveur"
            return f"Je lance le serveur {framework} pour {project.label}."
        except Exception as exc:
            return f"Impossible de lancer le serveur pour {project.label} : {exc}"

    def _run_git(self, args, project_key=None):
        project = self._resolve_project(project_key)
        if project is None:
            return "Choisis d'abord un projet Git."
        if not (project.path / ".git").exists():
            return f"{project.label} n'est pas un depot Git detectable."

        try:
            result = subprocess.run(
                ["git", *args],
                cwd=project.path,
                capture_output=True,
                text=True,
                timeout=config.GIT_TIMEOUT_SECONDS,
            )
        except Exception as exc:
            return f"Commande Git impossible sur {project.label} : {exc}"

        output = (result.stdout or result.stderr).strip()
        if not output:
            output = "Aucune sortie."
        return output

    def git_status(self, project_key=None):
        return self._run_git(["status", "--short", "--branch"], project_key=project_key)

    def git_pull(self, project_key=None):
        return self._run_git(["pull"], project_key=project_key)

    def git_log(self, project_key=None):
        return self._run_git(["log", "--oneline", "-5"], project_key=project_key)

    def prepare_commit(self, message="feat: update project", project_key=None):
        project = self._resolve_project(project_key)
        if project is None:
            return "Choisis d'abord un projet avant de preparer un commit."
        return (
            f"Commit prepare pour {project.label}. Confirmation obligatoire.\n"
            f"Commande suggeree : git add . && git commit -m \"{message}\""
        )

    def prepare_push(self, project_key=None):
        project = self._resolve_project(project_key)
        if project is None:
            return "Choisis d'abord un projet avant de preparer un push."
        return (
            f"Push prepare pour {project.label}. Confirmation obligatoire.\n"
            "Commande suggeree : git push"
        )

    def maintenance_checklist(self, project_key=None):
        project = self._resolve_project(project_key)
        label = project.label if project else "le projet actif"
        return (
            f"Checklist maintenance pour {label}:\n"
            "- verifier git status\n"
            "- lancer les tests du projet\n"
            "- verifier les variables d'environnement\n"
            "- relire les logs recents\n"
            "- confirmer les sauvegardes et la documentation"
        )

    def summarize_tabular_file(self, file_path):
        path = Path(file_path)
        if not path.exists():
            return f"Le fichier {path} est introuvable."

        if path.suffix.lower() == ".csv":
            with path.open("r", encoding="utf-8", newline="") as handle:
                rows = list(csv.reader(handle))
            preview = rows[:4]
            return "\n".join([", ".join(row) for row in preview]) or "CSV vide."

        if path.suffix.lower() == ".xlsx":
            try:
                from openpyxl import load_workbook
            except Exception as exc:
                return f"Lecture Excel indisponible : {exc}"

            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = workbook.active
            lines = []
            for row in sheet.iter_rows(max_row=4, values_only=True):
                lines.append(", ".join("" if cell is None else str(cell) for cell in row))
            workbook.close()
            return "\n".join(lines) or "Excel vide."

        return "Je supporte seulement CSV et XLSX pour le moment."
